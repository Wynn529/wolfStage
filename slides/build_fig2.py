"""Build the editable six-player state-flow diagram and a report PNG."""

from __future__ import annotations

from html import escape
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DIAGRAMS = ROOT / "slides" / "assets" / "diagrams"
FIGURES = ROOT / "report" / "figures"
DIAGRAMS.mkdir(parents=True, exist_ok=True)
FIGURES.mkdir(parents=True, exist_ok=True)
SVG = DIAGRAMS / "fig2_state_flow.svg"
PPTX = DIAGRAMS / "fig2_state_flow.pptx"
PREVIEW = DIAGRAMS / "fig2_state_flow_preview.png"
REPORT = FIGURES / "fig2_state_flow.png"

W, H = 1920, 1080
NAVY = "#0A1528"
PANEL = "#14243D"
CARD = "#203451"
WHITE = "#F3F7FC"
MUTED = "#BAC9D8"
CYAN = "#41D4D7"
AMBER = "#FFBE62"
GRAY = "#A8B7C9"
RED = "#F18B8B"
FONT_NAME = "Microsoft YaHei"
FONT_PATH = Path(r"C:\Windows\Fonts\msyh.ttc")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
SX, SY = 13.333 / W, 7.5 / H
canvas = Image.new("RGB", (W, H), NAVY)
draw = ImageDraw.Draw(canvas)
svg: list[str] = [
    f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">',
    '<defs>' + ''.join(
        f'<marker id="arrow-{name}" markerWidth="10" markerHeight="10" refX="9" refY="5" orient="auto">'
        f'<path d="M0,0 L10,5 L0,10 Z" fill="{color}"/></marker>'
        for name, color in (("cyan", CYAN), ("amber", AMBER), ("gray", GRAY))
    ) + '</defs>',
]


def rgb(color: str) -> RGBColor:
    return RGBColor.from_string(color.removeprefix("#"))


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_PATH), size, index=0)


def text(x: int, y: int, value: str, size: int = 24, color: str = WHITE,
         bold: bool = False, center: bool = False) -> None:
    anchor = "middle" if center else "start"
    svg.append(f'<text x="{x}" y="{y}" text-anchor="{anchor}" font-family="{FONT_NAME},sans-serif" '
               f'font-size="{size}" font-weight="{700 if bold else 400}" fill="{color}">{escape(value)}</text>')
    f = font(size, bold)
    bbox = draw.textbbox((0, 0), value, font=f)
    tx = x - (bbox[2] - bbox[0]) // 2 if center else x
    draw.text((tx, y - size), value, font=f, fill=color)
    width = max(60, min(1700, (bbox[2] - bbox[0]) + 12))
    bx = x - width / 2 if center else x
    box = slide.shapes.add_textbox(Inches(bx * SX), Inches((y - size - 3) * SY),
                                   Inches(width * SX), Inches(size * 1.45 * SY))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = value
    p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
    p.font.name = FONT_NAME
    p.font.size = Pt(size * .69)
    p.font.bold = bold
    p.font.color.rgb = rgb(color)


def rect(x: int, y: int, w: int, h: int, fill: str, stroke: str,
         radius: int = 14, stroke_width: int = 2) -> None:
    svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{radius}" '
               f'fill="{fill}" stroke="{stroke}" stroke-width="{stroke_width}"/>')
    draw.rounded_rectangle((x, y, x + w, y + h), radius, fill=fill, outline=stroke, width=stroke_width)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x * SX), Inches(y * SY),
                                   Inches(w * SX), Inches(h * SY))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(stroke)
    shape.line.width = Pt(stroke_width * .65)
    if shape.adjustments:
        shape.adjustments[0] = min(.2, radius / min(w, h))


def card(x: int, y: int, w: int, h: int, lines: list[str], color: str = CYAN,
         fill: str = CARD) -> None:
    rect(x, y, w, h, fill, color)
    step = 28 if len(lines) > 1 else 0
    start = y + h / 2 + 8
    for i, line in enumerate(lines):
        text(x + w // 2, int(start + (i - (len(lines) - 1) / 2) * step), line,
             23 if i == 0 else 20, WHITE if i == 0 else MUTED, i == 0, True)


def diamond(cx: int, cy: int, w: int, h: int, lines: list[str], color: str = CYAN) -> None:
    points = [(cx, cy-h//2), (cx+w//2, cy), (cx, cy+h//2), (cx-w//2, cy)]
    svg.append(f'<polygon points="{" ".join(f"{x},{y}" for x,y in points)}" '
               f'fill="{CARD}" stroke="{color}" stroke-width="2"/>')
    draw.polygon(points, fill=CARD)
    draw.line(points + [points[0]], fill=color, width=2)
    s = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches((cx-w/2)*SX), Inches((cy-h/2)*SY),
                               Inches(w*SX), Inches(h*SY))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(CARD)
    s.line.color.rgb = rgb(color); s.line.width = Pt(1.4)
    for i, line in enumerate(lines):
        text(cx, cy + 4 + (i - (len(lines)-1)/2) * 24, line, 19, WHITE, True, True)


def path(points: list[tuple[int, int]], color: str = GRAY, arrow: bool = True,
         width: int = 3) -> None:
    pts = " ".join(f"{x},{y}" for x, y in points)
    marker_name = {CYAN: "cyan", AMBER: "amber", GRAY: "gray"}[color]
    svg.append(f'<polyline points="{pts}" fill="none" stroke="{color}" stroke-width="{width}" '
               f'stroke-linejoin="round" stroke-linecap="round" '
               f'{f"marker-end=\"url(#arrow-{marker_name})\"" if arrow else ""}/>')
    draw.line(points, fill=color, width=width, joint="curve")
    segments = list(zip(points, points[1:]))
    for i, ((x1,y1),(x2,y2)) in enumerate(segments):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1*SX), Inches(y1*SY),
                                       Inches(x2*SX), Inches(y2*SY))
        c.line.color.rgb = rgb(color); c.line.width = Pt(width*.7)
        if arrow and i == len(segments)-1:
            tail = OxmlElement('a:tailEnd')
            tail.set('type', 'triangle')
            tail.set('w', 'med')
            tail.set('len', 'med')
            c.line._get_or_add_ln().append(tail)
    if arrow:
        import math
        x1,y1 = points[-2]; x2,y2 = points[-1]
        angle = math.atan2(y2-y1, x2-x1)
        left = (int(x2-12*math.cos(angle-.5)), int(y2-12*math.sin(angle-.5)))
        right = (int(x2-12*math.cos(angle+.5)), int(y2-12*math.sin(angle+.5)))
        draw.polygon([(x2,y2),left,right], fill=color)


# Background and section bands.
rect(0, 0, W, H, NAVY, NAVY, 0, 0)
text(64, 72, "固定六人局的状态转移", 45, WHITE, True)
text(1690, 70, "规则设计示意 · rules-v1.0", 20, MUTED, center=True)
rect(56, 118, 1808, 354, PANEL, "#314B68", 20)
rect(56, 490, 1808, 455, PANEL, "#314B68", 20)
text(82, 158, "01  夜晚：首夜开始，此后按同一流程循环", 24, AMBER, True)
text(82, 530, "02  白天：公开发言、分类、投票与放逐", 24, CYAN, True)

# Night: the two disagreement exits converge before inspection.
card(92, 205, 222, 70, ["两狼各自提案", "目标与自述理由"], AMBER)
diamond(400, 240, 132, 106, ["目标", "一致？"], AMBER)
card(515, 337, 215, 68, ["私密协商一次"], AMBER)
diamond(820, 371, 132, 106, ["协商后", "一致？"], AMBER)
card(920, 337, 218, 68, ["固定种子裁决", "从有效提案选一"], AMBER)
card(1190, 205, 210, 70, ["预言家查验", "私下获得结果"], AMBER)
card(1450, 205, 205, 70, ["夜间结算", "确定死亡玩家"], GRAY)
diamond(1750, 240, 130, 106, ["胜负", "已定？"], GRAY)

path([(314,240),(334,240)], AMBER)
path([(400,187),(400,181),(1295,181),(1295,205)], AMBER)
text(463, 178, "是", 18, AMBER, True)
path([(400,293),(400,371),(515,371)], AMBER)
text(421, 331, "否", 18, AMBER, True)
path([(730,371),(754,371)], AMBER)
path([(820,318),(820,295),(1260,295),(1260,275)], AMBER)
text(858, 306, "是", 18, AMBER, True)
path([(886,371),(920,371)], AMBER)
text(894, 359, "否", 18, AMBER, True)
path([(1138,371),(1165,371),(1165,240),(1190,240)], AMBER)
path([(1400,240),(1450,240)], GRAY)
path([(1655,240),(1685,240)], GRAY)

# Night resolution connects to the daytime entrance and the common terminal.
path([(1750,293),(1750,455),(152,455),(152,560)], GRAY)
text(1607, 446, "否 → 天亮", 19, GRAY, True)
path([(1815,240),(1882,240),(1882,1008),(1850,1008)], GRAY)
text(1795, 330, "是", 18, GRAY, True)

# Day's main left-to-right chain.
card(84, 560, 168, 68, ["天亮播报", "仅死亡编号"], GRAY)
card(315, 560, 214, 68, ["存活者依编号", "各公开发言一次"], CYAN)
card(589, 560, 233, 68, ["小模型逐句分类", "标签供后续参考"], CYAN)
card(882, 560, 188, 68, ["存活者投票", "各投其他存活者"], CYAN)
card(1130, 560, 190, 68, ["记录并计票", "无效票不计数"], GRAY)
diamond(1415, 594, 126, 106, ["有有效", "票？"], GRAY)
diamond(1570, 594, 126, 106, ["最高票", "平票？"], GRAY)
card(1700, 560, 145, 68, ["放逐", "最高票者"], GRAY)
path([(252,594),(315,594)], CYAN)
path([(529,594),(589,594)], CYAN)
path([(822,594),(882,594)], CYAN)
path([(1070,594),(1130,594)], GRAY)
path([(1320,594),(1352,594)], GRAY)
path([(1478,594),(1507,594)], GRAY)
text(1481, 579, "是", 18, GRAY, True)
path([(1633,594),(1700,594)], GRAY)
text(1654, 579, "否", 18, GRAY, True)

# Tie: one restricted revote, then either exile or no exile.
card(1390, 710, 225, 70, ["只对平票者", "重投一次"], GRAY)
diamond(1190, 838, 132, 108, ["仍然", "平票？"], GRAY)
card(950, 805, 158, 66, ["无人放逐"], GRAY)
path([(1415,647),(1415,677),(1029,677),(1029,805)], GRAY)
text(1428, 673, "否", 18, GRAY, True)
path([(1570,647),(1570,690),(1500,690),(1500,710)], GRAY)
text(1580, 677, "是", 18, GRAY, True)
path([(1390,745),(1190,745),(1190,784)], GRAY)
path([(1124,838),(1108,838)], GRAY)
text(1115, 824, "是", 18, GRAY, True)
path([(1256,838),(1772,838),(1772,628)], GRAY)
text(1290, 824, "否", 18, GRAY, True)

# Both outcomes are checked under the same deterministic rule.
diamond(1550, 884, 130, 100, ["胜负", "已定？"], GRAY)
path([(1772,628),(1772,797),(1550,797),(1550,834)], GRAY)
path([(1029,871),(1029,922),(1485,922)], GRAY)
diamond(1790, 884, 125, 100, ["已过", "8 昼夜？"], GRAY)
path([(1615,884),(1728,884)], GRAY)
text(1660, 869, "否", 18, GRAY, True)

# End states and loop, visually separated from the game actions.
rect(56, 965, 1808, 88, PANEL, "#314B68", 18)
card(102, 978, 195, 60, ["下一夜"], AMBER)
card(1050, 978, 226, 60, ["未完成终局"], GRAY)
card(1640, 978, 210, 60, ["胜方终局"], CYAN)
path([(102,1008),(29,1008),(29,240),(92,240)], AMBER)
path([(1550,934),(1550,1008),(1640,1008)], CYAN)
text(1570, 954, "是", 18, CYAN, True)
path([(1790,934),(1790,957),(1163,957),(1163,978)], GRAY)
text(1750, 950, "是", 18, GRAY, True)
path([(1728,884),(1690,884),(1690,948),(297,948),(297,1008)], AMBER)
text(1650, 940, "否", 18, AMBER, True)

# Compact rule notes belong to the graphic but do not claim implementation status.
text(84, 677, "旁注：预言家同夜被击杀，仍完成当夜查验。", 20, MUTED)
text(84, 711, "非法输出：重试一次；再次失败则暂停并记录错误。", 20, MUTED)
text(84, 745, "主持仅播报规则裁决后的事实，不决定胜负。", 20, MUTED)
text(84, 779, "无有效票：无人放逐；达到 8 昼夜无胜方：记未完成。", 20, MUTED)

svg.append("</svg>")
SVG.write_text("\n".join(svg), encoding="utf-8")
prs.save(PPTX)
canvas.save(PREVIEW)
canvas.save(REPORT, dpi=(300,300))
for p in (SVG, PPTX, PREVIEW, REPORT):
    print(p.relative_to(ROOT))
