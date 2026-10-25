"""Build the editable system architecture diagram for slides and reports."""

from __future__ import annotations

from html import escape
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "slides" / "assets" / "diagrams"
OUT.mkdir(parents=True, exist_ok=True)
SVG = OUT / "fig1_architecture.svg"
PPTX = OUT / "fig1_architecture.pptx"

W, H = 1920, 1080
NAVY = "#0A1528"
PANEL = "#14243D"
INNER = "#1D304D"
WHITE = "#F3F7FC"
MUTED = "#B9C7D9"
CYAN = "#41D4D7"
AMBER = "#FFBE62"
GRAY = "#A8B7C9"
FONT = "Microsoft YaHei"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#"))


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
SX, SY = 13.333 / W, 7.5 / H
svg: list[str] = [
    f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">',
    '<defs><marker id="arrow-cyan" markerWidth="12" markerHeight="12" refX="10" refY="6" orient="auto"><path d="M0,0 L12,6 L0,12 Z" fill="#41D4D7"/></marker>',
    '<marker id="arrow-amber" markerWidth="12" markerHeight="12" refX="10" refY="6" orient="auto"><path d="M0,0 L12,6 L0,12 Z" fill="#FFBE62"/></marker>',
    '<marker id="arrow-gray" markerWidth="12" markerHeight="12" refX="10" refY="6" orient="auto"><path d="M0,0 L12,6 L0,12 Z" fill="#A8B7C9"/></marker></defs>',
]


def rect(x, y, w, h, fill, stroke=None, radius=22, line=2):
    stroke = stroke or fill
    svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{radius}" fill="{fill}" stroke="{stroke}" stroke-width="{line}"/>')
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x * SX), Inches(y * SY), Inches(w * SX), Inches(h * SY))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    s.line.color.rgb = rgb(stroke); s.line.width = Pt(line * .65)
    if s.adjustments: s.adjustments[0] = min(.2, radius / min(w, h))


def txt(x, y, value, size=27, color=WHITE, bold=False, anchor="start"):
    svg.append(f'<text x="{x}" y="{y}" font-family="{FONT}, sans-serif" font-size="{size}" font-weight="{700 if bold else 400}" fill="{color}" text-anchor="{anchor}">{escape(value)}</text>')
    lines = value.split("\n")
    width = 1300 if anchor == "start" else 440
    box_x = x if anchor == "start" else x - width / 2
    t = slide.shapes.add_textbox(Inches(box_x * SX), Inches((y - size + 2) * SY), Inches(width * SX), Inches((size * 1.35 * len(lines)) * SY))
    tf = t.text_frame; tf.clear(); tf.margin_left = tf.margin_right = 0; tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = FONT; p.font.size = Pt(size * .73); p.font.bold = bold; p.font.color.rgb = rgb(color)
        p.alignment = PP_ALIGN.LEFT if anchor == "start" else PP_ALIGN.CENTER


def line(x1, y1, x2, y2, color, dashed=False, arrow=True, width=4):
    dash = ' stroke-dasharray="10 8"' if dashed else ""
    marker_id = {CYAN: "cyan", AMBER: "amber", GRAY: "gray"}[color]
    svg.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="{color}" stroke-width="{width}"{dash}{f" marker-end=\"url(#arrow-{marker_id})\"" if arrow else ""}/>')
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1 * SX), Inches(y1 * SY), Inches(x2 * SX), Inches(y2 * SY))
    c.line.color.rgb = rgb(color); c.line.width = Pt(width * .7)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH


rect(0, 0, W, H, NAVY, NAVY, 0, 0)
txt(80, 75, "系统架构与信息流", 45, WHITE, True)
txt(1840, 73, "系统设计示意", 23, MUTED, False, "middle")

rect(80, 125, 1760, 130, PANEL, CYAN)
txt(115, 170, "上帝视角网页界面", 31, WHITE, True)
txt(115, 212, "六张玩家卡 · 阶段与事件时间线 · 自述理由 · 分类标签 · 调用统计", 24, MUTED)
txt(1640, 171, "自动运行 / 暂停 / 单步", 24, CYAN, True, "middle")
txt(1640, 212, "只读回放 / 日志导出", 23, MUTED, False, "middle")
line(885, 257, 885, 326, CYAN)
line(1035, 326, 1035, 257, GRAY)
txt(735, 299, "控制请求", 21, CYAN)
txt(1062, 299, "查询 / 回放", 21, GRAY)

rect(80, 327, 1760, 324, PANEL, GRAY)
txt(115, 373, "FastAPI 游戏服务", 32, WHITE, True)
txt(1650, 372, "主持全局真相与权限", 23, MUTED, False, "middle")

cards = [
    (110, "事件推进", ["创建对局 / 昼夜阶段", "自动运行与事件边界暂停"]),
    (545, "角色上下文", ["公开事件 + 本人可见私密事件", "按 agent_id 过滤输入"]),
    (980, "确定性规则", ["行动校验 / 计票 / 胜负", "结构化行动为正式依据"]),
    (1415, "事件与调用记录", ["events.jsonl / calls.jsonl", "状态快照 / 统计 / 原始回放"]),
]
for x, heading, body in cards:
    rect(x, 420, 395, 164, INNER, "#38516F", 18)
    txt(x + 24, 463, heading, 28, WHITE, True)
    txt(x + 24, 509, body[0], 20, MUTED)
    txt(x + 24, 547, body[1], 20, MUTED)
line(505, 502, 530, 502, GRAY)
line(940, 502, 965, 502, GRAY)
line(1375, 502, 1400, 502, GRAY)
txt(115, 626, "公开发言与分类结果写入同一事件链；规则程序独立裁决", 22, GRAY)

rect(80, 757, 760, 235, PANEL, AMBER)
txt(112, 803, "六个角色代理", 29, WHITE, True)
txt(112, 839, "2 狼人 · 1 预言家 · 3 平民", 23, AMBER)
txt(112, 878, "共享 DeepSeek 基础模型；独立提示词、可见历史和私密记忆", 21, MUTED)
txt(112, 916, "私密：身份 / 狼队提案 / 查验结果 / 本人自述理由", 20, MUTED)
txt(112, 954, "输出：发言、击杀、查验、投票的结构化行动", 20, MUTED)

rect(875, 757, 345, 235, PANEL, AMBER)
txt(902, 803, "主持语言代理", 28, WHITE, True)
txt(902, 846, "仅接收已判定事实", 21, MUTED)
txt(902, 884, "阶段播报 / 提醒", 21, MUTED)
txt(902, 922, "总结与赛后复盘", 21, MUTED)
txt(902, 960, "不修改规则事实", 20, AMBER)

rect(1255, 757, 585, 235, PANEL, CYAN)
txt(1283, 803, "A：公开发言分析接口", 28, WHITE, True)
txt(1283, 842, "输入 game_id / event_id / speaker_id / public_text", 20, CYAN)
txt(1283, 879, "在线：意图分类 + 独立规则实体抽取", 21, MUTED)
txt(1283, 916, "输出 intent / 未校准 score / 实体 / 版本", 20, MUTED)
txt(1283, 953, "离线：数据训练与冻结评估", 20, MUTED)

line(360, 680, 360, 755, AMBER)
line(500, 755, 500, 680, AMBER)
txt(115, 713, "授权上下文", 19, AMBER)
txt(512, 713, "结构化行动", 19, AMBER)
line(1015, 680, 1015, 755, AMBER)
line(1135, 755, 1135, 680, AMBER)
txt(887, 713, "判定事实", 19, AMBER)
txt(1142, 713, "主持措辞", 19, AMBER)
line(1450, 680, 1450, 755, CYAN)
line(1670, 755, 1670, 680, CYAN)
txt(1280, 713, "公开发言", 19, CYAN)
txt(1680, 713, "标签与实体", 19, CYAN)

line(104, 1031, 180, 1031, CYAN, arrow=False)
txt(194, 1039, "公开信息", 19, MUTED)
line(385, 1031, 461, 1031, AMBER, arrow=False)
txt(475, 1039, "私密 / 授权信息", 19, MUTED)
line(755, 1031, 831, 1031, GRAY, arrow=False)
txt(845, 1039, "规则、日志与只读查询", 19, MUTED)
txt(1630, 1039, "上帝视角内容不回流角色", 19, MUTED, False, "middle")

svg.append("</svg>")
SVG.write_text("\n".join(svg), encoding="utf-8")
prs.save(PPTX)
print(SVG.relative_to(ROOT))
print(PPTX.relative_to(ROOT))
